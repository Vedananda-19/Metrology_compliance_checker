from datetime import datetime, timezone
from io import BytesIO

from services.report import rules_context, justify
from pipeline.extraction.declarations import FACT_MAP
from pipeline.measurement import font as font_measure

FONT_LABELS = {field: label for field, (_, _, label) in FACT_MAP.items()}

DECISION_LABEL = {
    "CONFIRMED": "Violation confirmed by officer",
    "DISMISSED": "Dismissed as a false reading",
    "VERIFIED_COMPLIANT": "Verified compliant on inspection",
    "VERIFIED_NON_COMPLIANT": "Verified as a violation on inspection",
}

CONFIRMED_DECISIONS = {"CONFIRMED", "VERIFIED_NON_COMPLIANT"}
CLEARED_DECISIONS = {"DISMISSED", "VERIFIED_COMPLIANT"}


def _person(user) -> str:
    if user is None:
        return "—"
    return user.full_name or user.username


def _finding_row(finding: dict, review: dict | None) -> dict:
    provision = rules_context.provision_for(finding["rule_id"]) or {}
    decision = review.get("decision") if review else None
    return {
        "rule_id": finding["rule_id"],
        "rule_ref": finding.get("rule_ref"),
        "rule_number": provision.get("rule") or finding.get("rule_ref"),
        "requirement": finding.get("requirement"),
        "severity": finding.get("severity"),
        "status": finding["status"],
        "observed": finding.get("observed"),
        "explanation": finding.get("explanation"),
        "official_text": provision.get("text", ""),
        "pdf_page": provision.get("pdf_page"),
        "decision": decision,
        "decision_label": DECISION_LABEL.get(decision) if decision else None,
        "review_note": review.get("note") if review else None,
    }


def gather(db, inspection) -> dict:
    result = (inspection.evaluation.result if inspection.evaluation else {}) or {}
    findings = result.get("findings", [])
    reviews = {review.rule_id: {"decision": review.decision, "note": review.note} for review in inspection.reviews}

    confirmed, cleared, not_verifiable = [], [], []
    for finding in findings:
        if finding["status"] not in ("NON_COMPLIANT", "NOT_VERIFIABLE"):
            continue
        review = reviews.get(finding["rule_id"])
        decision = review.get("decision") if review else None
        row = _finding_row(finding, review)
        if decision in CLEARED_DECISIONS:
            cleared.append(row)
        elif decision in CONFIRMED_DECISIONS or finding["status"] == "NON_COMPLIANT":
            confirmed.append(row)
        else:
            not_verifiable.append(row)

    written = justify.justify(
        [
            {
                "rule_id": row["rule_id"],
                "rule_ref": row["rule_ref"],
                "requirement": row["requirement"],
                "explanation": row["explanation"],
                "observed": row["observed"],
            }
            for row in confirmed
        ]
    )
    for row in confirmed:
        row["why"] = written.get(row["rule_id"]) or row["explanation"]

    fonts = []
    for measurement in inspection.font_measurements:
        fonts.append({
            "field": measurement.field,
            "label": FONT_LABELS.get(measurement.field, measurement.field),
            "measured_height_mm": measurement.measured_height_mm,
            "reference_size_mm": measurement.reference_size_mm,
            "status": measurement.status,
            "detail": measurement.detail,
        })
    fonts.sort(key=lambda row: (row["status"] != font_measure.MEASURED, row["label"]))

    counts = result.get("counts", {})
    return {
        "reference": inspection.reference,
        "fonts": fonts,
        "title": inspection.title or "Untitled package",
        "generated_at": datetime.now(timezone.utc),
        "verified_at": inspection.verified_at,
        "created_by": _person(inspection.inspector),
        "assigned_to": _person(inspection.assignee) if inspection.assignee else None,
        "verified_by": _person(inspection.verifier) if inspection.verifier else None,
        "engine_verdict": result.get("verdict"),
        "rule_set_version": result.get("rule_set_version"),
        "summary": result.get("summary"),
        "engine_counts": counts,
        "confirmed": confirmed,
        "cleared": cleared,
        "not_verifiable": not_verifiable,
        "totals": {
            "confirmed_violations": len(confirmed),
            "cleared": len(cleared),
            "not_verifiable": len(not_verifiable),
            "compliant": counts.get("COMPLIANT", 0),
            "not_applicable": counts.get("NOT_APPLICABLE", 0),
            "exempt": counts.get("EXEMPT", 0),
        },
    }


def _fmt(when) -> str:
    if when is None:
        return "—"
    return when.strftime("%d %b %Y, %H:%M UTC")


def render_pdf(data: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable,
    )

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=9.5, leading=13)
    small = ParagraphStyle("small", parent=body, fontSize=8, textColor=colors.HexColor("#5b6b7f"))
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=17, textColor=colors.HexColor("#0e2846"))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12.5, textColor=colors.HexColor("#0e2846"), spaceBefore=6)
    rule_head = ParagraphStyle("rule", parent=body, fontSize=10.5, textColor=colors.HexColor("#c02626"), spaceBefore=2)
    quote = ParagraphStyle("quote", parent=body, fontSize=9, leftIndent=8, textColor=colors.HexColor("#33405a"))
    why_style = ParagraphStyle("why", parent=body, fontSize=8.7, leading=11.5)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=16 * mm, bottomMargin=16 * mm, leftMargin=18 * mm, rightMargin=18 * mm)
    flow = []

    flow.append(Paragraph("MetroGuard — Legal Metrology Compliance Report", h1))
    flow.append(Paragraph("Legal Metrology (Packaged Commodities) Rules, 2011", small))
    flow.append(Spacer(1, 6))
    flow.append(HRFlowable(width="100%", color=colors.HexColor("#e2e8f1")))
    flow.append(Spacer(1, 8))

    meta = [
        ["Inspection", data["reference"], "Package", data["title"]],
        ["Field officer", data["created_by"], "Assigned to", data["assigned_to"] or "—"],
        ["Verified by", data["verified_by"] or "—", "Verified at", _fmt(data["verified_at"])],
        ["Report generated", _fmt(data["generated_at"]), "Rule set", data["rule_set_version"] or "—"],
    ]
    table = Table(meta, colWidths=[28 * mm, 60 * mm, 26 * mm, 56 * mm])
    table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#5b6b7f")),
        ("TEXTCOLOR", (2, 0), (2, -1), colors.HexColor("#5b6b7f")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    flow.append(table)
    flow.append(Spacer(1, 10))

    totals = data["totals"]
    flow.append(Paragraph("Compliance summary", h2))
    summary_rows = [
        ["Confirmed violations", "Cleared on review", "Compliant", "Exempt", "Not applicable"],
        [
            str(totals["confirmed_violations"]),
            str(totals["cleared"]),
            str(totals["compliant"]),
            str(totals["exempt"]),
            str(totals["not_applicable"]),
        ],
    ]
    summary = Table(summary_rows, colWidths=[34 * mm] * 5)
    summary.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("FONTSIZE", (0, 1), (-1, 1), 15),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#5b6b7f")),
        ("TEXTCOLOR", (0, 1), (0, 1), colors.HexColor("#c02626")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 1), (-1, 1), 2),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f2f5f9")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f1")),
    ]))
    flow.append(summary)
    if data["summary"]:
        flow.append(Spacer(1, 4))
        flow.append(Paragraph(data["summary"], small))
    flow.append(Spacer(1, 12))

    flow.append(Paragraph("Violations (%d)" % totals["confirmed_violations"], h2))
    if not data["confirmed"]:
        flow.append(Paragraph("No violations were confirmed for this package.", body))
    for index, row in enumerate(data["confirmed"], start=1):
        flow.append(Spacer(1, 6))
        flow.append(Paragraph("%d. Rule %s — %s" % (index, row["rule_number"], row["requirement"]), rule_head))
        tag = ("  ·  " + row["decision_label"]) if row["decision_label"] else ""
        flow.append(Paragraph("Severity: %s%s" % (row["severity"], tag), small))
        if row["observed"]:
            flow.append(Paragraph("<b>Observed on package:</b> %s" % row["observed"], body))
        if row["why"]:
            flow.append(Paragraph("<b>Why this is a violation:</b> %s" % row["why"], why_style))
        if row["official_text"]:
            page = (" (source p.%s)" % row["pdf_page"]) if row["pdf_page"] else ""
            flow.append(Paragraph("<i>Rule text%s:</i> %s" % (page, row["official_text"]), quote))
        if row["review_note"]:
            flow.append(Paragraph("Officer note: %s" % row["review_note"], small))
        flow.append(HRFlowable(width="100%", color=colors.HexColor("#eef2f7")))

    if data["cleared"]:
        flow.append(Spacer(1, 10))
        flow.append(Paragraph("Cleared on review (%d)" % len(data["cleared"]), h2))
        for row in data["cleared"]:
            note = (". " + row["review_note"]) if row["review_note"] else ""
            flow.append(Paragraph(
                "Rule %s — %s: %s%s" % (row["rule_number"], row["requirement"], row["decision_label"], note),
                body,
            ))

    if data["not_verifiable"]:
        flow.append(Spacer(1, 10))
        flow.append(Paragraph("Could not be verified from the images (%d)" % len(data["not_verifiable"]), h2))
        flow.append(Paragraph(
            "These requirements need a physical measurement or a check the photographs cannot establish. "
            "They are neither passed nor failed.",
            small,
        ))
        for row in data["not_verifiable"]:
            flow.append(Spacer(1, 4))
            flow.append(Paragraph("Rule %s — %s" % (row["rule_number"], row["requirement"]), body))
            if row["observed"]:
                flow.append(Paragraph(row["observed"], small))
            if row["official_text"]:
                page = (" (source p.%s)" % row["pdf_page"]) if row["pdf_page"] else ""
                flow.append(Paragraph("<i>Rule text%s:</i> %s" % (page, row["official_text"]), quote))

    if data["fonts"]:
        flow.append(Spacer(1, 12))
        flow.append(Paragraph("Printed character height (Rule 7)", h2))
        font_rows = [["Declaration", "Measured height", "Reference marker", "Status"]]
        for row in data["fonts"]:
            height = "%.2f mm" % row["measured_height_mm"] if row["measured_height_mm"] is not None else "—"
            status = "Measured" if row["status"] == "MEASURED" else (row["detail"] or "Unable to verify")
            font_rows.append([row["label"], height, "%.0f mm" % row["reference_size_mm"], status])
        font_table = Table(font_rows, colWidths=[64 * mm, 30 * mm, 32 * mm, 44 * mm])
        font_table.setStyle(TableStyle([
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#5b6b7f")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f2f5f9")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e2e8f1")),
            ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#e2e8f1")),
        ]))
        flow.append(font_table)
        flow.append(Spacer(1, 3))
        flow.append(Paragraph(
            "Heights are measured from the package image using the ArUco reference marker for scale. "
            "Rule 7 sets the minimum height by principal-display-panel area.",
            small,
        ))

    flow.append(Spacer(1, 14))
    flow.append(HRFlowable(width="100%", color=colors.HexColor("#e2e8f1")))
    flow.append(Spacer(1, 4))
    flow.append(Paragraph(
        "Compliance was adjudicated by a deterministic rule engine and confirmed by the named officer. "
        "The machine-readable rules follow the Legal Metrology (Packaged Commodities) Rules, 2011.",
        small,
    ))

    doc.build(flow)
    return buffer.getvalue()


def render_docx(data: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    document = Document()
    document.add_heading("MetroGuard — Legal Metrology Compliance Report", level=0)
    document.add_paragraph("Legal Metrology (Packaged Commodities) Rules, 2011")

    meta = document.add_table(rows=0, cols=4)
    for label_a, value_a, label_b, value_b in [
        ("Inspection", data["reference"], "Package", data["title"]),
        ("Field officer", data["created_by"], "Assigned to", data["assigned_to"] or "—"),
        ("Verified by", data["verified_by"] or "—", "Verified at", _fmt(data["verified_at"])),
        ("Report generated", _fmt(data["generated_at"]), "Rule set", data["rule_set_version"] or "—"),
    ]:
        cells = meta.add_row().cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text = label_a, str(value_a), label_b, str(value_b)

    totals = data["totals"]
    document.add_heading("Compliance summary", level=1)
    summary = document.add_table(rows=2, cols=5)
    heads = ["Confirmed violations", "Cleared on review", "Compliant", "Exempt", "Not applicable"]
    vals = [
        totals["confirmed_violations"], totals["cleared"], totals["compliant"],
        totals["exempt"], totals["not_applicable"],
    ]
    for i, head in enumerate(heads):
        summary.rows[0].cells[i].text = head
        summary.rows[1].cells[i].text = str(vals[i])
    if data["summary"]:
        document.add_paragraph(data["summary"])

    document.add_heading("Violations (%d)" % totals["confirmed_violations"], level=1)
    if not data["confirmed"]:
        document.add_paragraph("No violations were confirmed for this package.")
    for index, row in enumerate(data["confirmed"], start=1):
        heading = document.add_heading("%d. Rule %s — %s" % (index, row["rule_number"], row["requirement"]), level=2)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0xC0, 0x26, 0x26)
        tag = ("  ·  " + row["decision_label"]) if row["decision_label"] else ""
        line = document.add_paragraph()
        run = line.add_run("Severity: %s%s" % (row["severity"], tag))
        run.italic = True
        run.font.size = Pt(9)
        if row["observed"]:
            p = document.add_paragraph()
            p.add_run("Observed on package: ").bold = True
            p.add_run(row["observed"])
        if row["why"]:
            p = document.add_paragraph()
            label = p.add_run("Why this is a violation: ")
            label.bold = True
            label.font.size = Pt(9)
            p.add_run(row["why"]).font.size = Pt(9)
        if row["official_text"]:
            page = (" (source p.%s)" % row["pdf_page"]) if row["pdf_page"] else ""
            p = document.add_paragraph()
            run = p.add_run("Rule text%s: %s" % (page, row["official_text"]))
            run.italic = True
        if row["review_note"]:
            document.add_paragraph("Officer note: %s" % row["review_note"])

    if data["cleared"]:
        document.add_heading("Cleared on review (%d)" % len(data["cleared"]), level=1)
        for row in data["cleared"]:
            note = (". " + row["review_note"]) if row["review_note"] else ""
            document.add_paragraph(
                "Rule %s — %s: %s%s" % (row["rule_number"], row["requirement"], row["decision_label"], note)
            )

    if data["not_verifiable"]:
        document.add_heading("Could not be verified from the images (%d)" % len(data["not_verifiable"]), level=1)
        document.add_paragraph(
            "These requirements need a physical measurement or a check the photographs cannot establish. "
            "They are neither passed nor failed."
        )
        for row in data["not_verifiable"]:
            p = document.add_paragraph()
            p.add_run("Rule %s — %s" % (row["rule_number"], row["requirement"])).bold = True
            if row["observed"]:
                document.add_paragraph(row["observed"])
            if row["official_text"]:
                page = (" (source p.%s)" % row["pdf_page"]) if row["pdf_page"] else ""
                run = document.add_paragraph().add_run("Rule text%s: %s" % (page, row["official_text"]))
                run.italic = True

    if data["fonts"]:
        document.add_heading("Printed character height (Rule 7)", level=1)
        table = document.add_table(rows=1, cols=4)
        headers = ["Declaration", "Measured height", "Reference marker", "Status"]
        for i, head in enumerate(headers):
            table.rows[0].cells[i].text = head
        for row in data["fonts"]:
            height = "%.2f mm" % row["measured_height_mm"] if row["measured_height_mm"] is not None else "—"
            status = "Measured" if row["status"] == "MEASURED" else (row["detail"] or "Unable to verify")
            cells = table.add_row().cells
            cells[0].text = row["label"]
            cells[1].text = height
            cells[2].text = "%.0f mm" % row["reference_size_mm"]
            cells[3].text = status
        document.add_paragraph(
            "Heights are measured from the package image using the ArUco reference marker for scale. "
            "Rule 7 sets the minimum height by principal-display-panel area."
        )

    document.add_paragraph(
        "Compliance was adjudicated by a deterministic rule engine and confirmed by the named officer. "
        "The machine-readable rules follow the Legal Metrology (Packaged Commodities) Rules, 2011."
    )

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
